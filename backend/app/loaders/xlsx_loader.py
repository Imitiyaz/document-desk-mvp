# SPDX-License-Identifier: AGPL-3.0-only
"""Excel XLSX/XLSM extraction with physical row/cell provenance.

The physical openpyxl path preserves native cell values, merge origins, formats
and formula-cache diagnostics. Cleaning and interpretation remain separate.
A DataFrame adapter is retained for caller compatibility.
See docs/pipeline.md and docs/data-model.md.
Author: Imitiyaz and AI Agent
Project: RAG MVP
Layer: Loaders
"""

from __future__ import annotations

from pathlib import Path
from typing import Any
from zipfile import ZipFile

import pandas as pd
from loguru import logger
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from app.core.config import settings
from app.core.exceptions import (
    CorruptFileError,
    ExtractionLimitError,
    LegacyExcelFormatError,
    UnsupportedLayoutError,
)
from app.loaders.base import BaseLoader
from app.models.extracted import ExtractedDocument, ExtractedRow, SourceMetadata

SUPPORTED_EXTENSIONS = {".xlsx", ".xlsm"}
LEGACY_EXTENSIONS = {".xls"}


class XlsxLoader(BaseLoader):
    """Loads `.xlsx`/`.xlsm` workbooks into `ExtractedDocument` records."""

    supported_extensions = frozenset(SUPPORTED_EXTENSIONS)

    def supports(self, file_path: Path) -> bool:
        # Deliberately includes legacy .xls too, so a .xls file routes here
        # and gets the specific, actionable LegacyExcelFormatError from
        # load() below -- rather than the generic "no loader found" error
        # from the registry in app/loaders/__init__.py.
        """Route supported suffixes; legacy XLS is routed to an explicit format error."""
        return file_path.suffix.lower() in SUPPORTED_EXTENSIONS | LEGACY_EXTENSIONS

    def load(self, file_path: Path) -> ExtractedDocument:
        """Extract physical cells with native types, merge origins and formula cache state.

        The physical openpyxl path avoids pandas coercing text identifiers/NA.
        The DataFrame adapter below remains for existing caller compatibility.
        Formula expressions are inspected only for diagnostics, never executed.
        """
        if file_path.suffix.lower() in LEGACY_EXTENSIONS:
            raise LegacyExcelFormatError(
                f"{file_path.name}: re-save legacy .xls format as .xlsx before uploading."
            )
        logger.info("Loading Excel workbook: {}", file_path.name)
        workbook = None
        formulas = None
        document = ExtractedDocument(source_file=file_path.name)
        try:
            with ZipFile(file_path) as archive:
                if (
                    sum(entry.file_size for entry in archive.infolist())
                    > settings.max_xlsx_expanded_bytes
                ):
                    raise ExtractionLimitError(
                        "Workbook expanded content exceeds the configured extraction budget."
                    )
            workbook = load_workbook(file_path, data_only=True, keep_links=False)
            formulas = load_workbook(
                file_path, data_only=False, read_only=True, keep_links=False
            )
            for sheet in workbook:
                if (
                    sheet.max_row > settings.max_extracted_rows + 1
                    or sheet.max_column > settings.max_xlsx_columns
                ):
                    raise ExtractionLimitError(
                        "Worksheet extent exceeds the configured row/column extraction budget."
                    )
                raw_headers = [
                    sheet.cell(1, col).value for col in range(1, sheet.max_column + 1)
                ]
                named = [
                    str(h).strip().casefold() for h in raw_headers if h is not None
                ]
                if len(named) != len(set(named)):
                    raise UnsupportedLayoutError(
                        f"Sheet '{sheet.title}' has duplicate column headers; export one clear header row."
                    )
                headers = [
                    str(h) if h is not None else f"column_{col}"
                    for col, h in enumerate(raw_headers, 1)
                ]
                formula_cells = {
                    cell.coordinate
                    for cells in formulas[sheet.title].iter_rows(min_row=2)
                    for cell in cells
                    if cell.data_type == "f"
                }
                for col, header in enumerate(raw_headers, 1):
                    if header is None and any(
                        sheet.cell(row, col).value is not None
                        or f"{get_column_letter(col)}{row}" in formula_cells
                        for row in range(2, sheet.max_row + 1)
                    ):
                        raise UnsupportedLayoutError(
                            f"Sheet '{sheet.title}' has values beneath a blank header; multi-row header reconstruction is not supported."
                        )
                anchors = {}
                for merged in sheet.merged_cells.ranges:
                    if merged.min_col == merged.max_col and merged.min_row > 1:
                        for row_number in range(merged.min_row, merged.max_row + 1):
                            anchors[(row_number, merged.min_col)] = sheet.cell(
                                merged.min_row, merged.min_col
                            )
                missing = []
                for number in range(2, sheet.max_row + 1):
                    data = {
                        h: sheet.cell(number, col).value
                        for col, h in enumerate(headers, 1)
                    }
                    row = ExtractedRow(
                        data=data,
                        metadata=SourceMetadata(file_path.name, sheet.title, number),
                    )
                    for col, header in enumerate(headers, 1):
                        cell = sheet.cell(number, col)
                        row.cell_origins[header] = cell.coordinate
                        row.cell_formats[header] = cell.number_format
                        if (number, col) in anchors:
                            cell = anchors[(number, col)]
                            row.merged_values[header] = cell.value
                            row.cell_origins[header] = cell.coordinate
                            row.cell_formats[header] = cell.number_format
                        if cell.coordinate in formula_cells:
                            row.formula_fields.add(header)
                            if cell.value is None:
                                missing.append(cell.coordinate)
                    if (
                        any(v not in (None, "") for v in data.values())
                        or row.formula_fields
                    ):
                        document.rows.append(row)
                    if len(document.rows) > settings.max_extracted_rows:
                        raise ExtractionLimitError(
                            "Workbook row count exceeds the extraction budget."
                        )
                if missing:
                    document.warnings.append(
                        f"{sheet.title}: {len(set(missing))} formula cell(s) lack cached values ({', '.join(sorted(set(missing))[:10])}). Recalculate/save in Excel, then re-upload. Formulas are not executed here."
                    )
        except (ExtractionLimitError, UnsupportedLayoutError):
            raise
        except Exception as exc:
            raise CorruptFileError(
                f"Failed to read Excel file {file_path.name}: {exc}"
            ) from exc
        finally:
            if workbook is not None:
                workbook.close()
            if formulas is not None:
                formulas.close()
        return document

    @staticmethod
    def _extract_sheet(
        df: pd.DataFrame,
        sheet_name: str,
        source_file: str,
        document: ExtractedDocument,
    ) -> None:
        """Append every non-empty row of one sheet to `document.rows`."""
        if df.empty:
            logger.warning(
                "Sheet '{}' in {} has no data, skipping", sheet_name, source_file
            )
            return

        headers = [str(col) for col in df.columns]

        for row_index, row in df.iterrows():
            # strict=True deliberately: a pandas DataFrame is
            # rectangular by construction -- every row has exactly as
            # many values as there are columns. A mismatch here would
            # mean something structurally wrong, worth failing loudly
            # on rather than silently truncating.
            row_data: dict[str, Any] = {
                header: (None if pd.isna(value) else value)
                for header, value in zip(headers, row.tolist(), strict=True)
            }

            if all(value is None for value in row_data.values()):
                continue  # entirely blank row -- not useful data, don't fabricate a chunk from it

            # pandas' row_index is 0-based and doesn't count the header row;
            # +2 converts it to the row number a user would actually see if
            # they opened this file in Excel (header = row 1, first data
            # row = row 2). This matters because a citation that says
            # "row 5" should mean row 5 in the actual spreadsheet.
            excel_row_number = int(row_index) + 2  # type: ignore[arg-type]

            document.rows.append(
                ExtractedRow(
                    data=row_data,
                    metadata=SourceMetadata(
                        source_file=source_file,
                        sheet_or_page=sheet_name,
                        row_or_location=excel_row_number,
                    ),
                )
            )
