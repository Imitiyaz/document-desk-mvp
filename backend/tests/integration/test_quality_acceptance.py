"""Independent source-cell and numeric oracles for the actual production API path.

Expected records/values below are reviewed fixture facts, not values obtained
from the application's loader, ranking, cleaner, parser or answer checker.

Author: Imitiyaz and AI Agent
Project: RAG MVP
Layer: Verification and synthetic fixtures
"""

from __future__ import annotations

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.main import create_app
from tests.api_fakes import make_bundle


@pytest.fixture
def api(fixtures_dir):
    bundle = make_bundle()
    with TestClient(create_app(lambda: bundle.service)) as client:
        for name in (
            "sample_sales.xlsx",
            "sample_sales_extended.xlsx",
            "sample_report.pdf",
            "sample_report_extended.pdf",
        ):
            response = client.post(
                "/documents", files={"file": (name, (fixtures_dir / name).read_bytes())}
            )
            assert response.status_code == 201, response.text
        yield client, bundle


def ask(api, question, source="sample_sales_extended.xlsx", **scope):
    response = api[0].post(
        "/query",
        json={"question": question, "source_file": source, "top_k": 1, **scope},
    )
    assert response.status_code == 200, response.text
    body = response.json()
    assert body["unresolved_citation_numbers"] == []
    if body["execution_path"] == "structured":
        assert api[1].llm.prompts == []  # structured arithmetic never uses a model
    return body


def identities(body, field="Q3 Revenue"):
    return [
        (r["sheet_or_page"], r["row_or_location"], r["values"].get(field))
        for r in body["records"]
    ]


def test_user_query_1_preserves_four_north_records_with_disclosed_continuation_even_at_top_k_one(
    api,
):
    b = ask(api, "What was the Q3 revenue for North?")
    assert b["status"] == "answered"
    assert identities(b) == [
        ("Merged Region", 2, "45000"),
        ("Merged Region", 3, "47000"),
        ("Sparse Data", 2, "45000"),
        ("Sparse Data", 3, "38000"),
    ]
    assert [c["row_or_location"] for c in b["citations"]] == [2, 3, 2, 3]
    assert [r["cell_references"]["Q3 Revenue"] for r in b["records"]] == [
        "B2",
        "B3",
        "B2",
        "B3",
    ]
    assert b["records"][1]["cell_references"]["Region"] == "A2"
    assert b["coverage"]["matched_records"] == b["coverage"]["returned_records"] == 4
    assert b["calculation"] is None


def test_user_query_2_lists_all_records_without_inventing_a_total(api):
    b = ask(api, "What was the Q3 revenue?")
    assert b["status"] == "incomplete"
    assert b["calculation"] is None
    assert identities(b) == [
        ("Merged Region", 2, "45000"),
        ("Merged Region", 3, "47000"),
        ("Merged Region", 4, "38000"),
        ("Merged Region", 5, "39000"),
        ("Sparse Data", 2, "45000"),
        ("Sparse Data", 3, "38000"),
        ("Sparse Data", 4, None),
        ("Sparse Data", 5, "41000"),
        ("Sparse Data", 6, "22000"),
    ]
    assert len(b["citations"]) == 9
    assert "140100" not in b["answer"]


def test_explicit_ambiguous_cross_sheet_total_requires_business_scope(api):
    b = ask(api, "Total Q3 revenue")
    assert b["status"] == "clarification_required"
    assert b["calculation"] is None
    assert set(b["coverage"]["available_scopes"]) == {"Merged Region", "Sparse Data"}


def test_scoped_sparse_records_preserve_blanks_and_disclose_continuation(api):
    b = ask(api, "Show Q3 revenue", sheet_or_page="Sparse Data")
    assert identities(b) == [
        ("Sparse Data", 2, "45000"),
        ("Sparse Data", 3, "38000"),
        ("Sparse Data", 4, None),
        ("Sparse Data", 5, "41000"),
        ("Sparse Data", 6, "22000"),
    ]
    assert [r["values"]["Region"] for r in b["records"]] == [
        "North",
        "North",
        "East",
        "West",
        "West",
    ]
    assert b["status"] == "incomplete"
    assert b["coverage"]["missing_cells"] == [
        {"number": 3, "field": "Q3 Revenue", "cell": "B4"}
    ]


def test_all_merged_source_records_retained_in_order(api):
    b = ask(api, "Show Q3 revenue", sheet_or_page="Merged Region")
    assert identities(b) == [
        ("Merged Region", 2, "45000"),
        ("Merged Region", 3, "47000"),
        ("Merged Region", 4, "38000"),
        ("Merged Region", 5, "39000"),
    ]


def test_scoped_total_has_complete_operand_provenance_and_independent_exact_sum(api):
    b = ask(api, "Total Q3 revenue in Merged Region")
    assert b["calculation"]["result"] == {"Q3 Revenue": "169000"}
    assert [o["values"]["Q3 Revenue"] for o in b["calculation"]["operands"]] == [
        "45000",
        "47000",
        "38000",
        "39000",
    ]
    assert b["coverage"]["complete"] is True
    assert len(b["citations"]) == 4


def test_partial_aggregate_never_treats_missing_cell_as_zero(api):
    b = ask(api, "Total Q3 revenue in Sparse Data")
    assert b["status"] == "incomplete"
    assert b["calculation"]["result"] == {"Q3 Revenue": "146000"}
    assert [o["number"] for o in b["calculation"]["operands"]] == [1, 2, 4, 5]
    assert b["calculation"]["excluded_cells"][0]["cell"] == "B4"
    assert len(b["records"]) == 5


def test_known_three_region_sum_correct_with_all_three_distinct_citations(api):
    b = ask(
        api,
        "What is the combined Q3 revenue for North, South, and East in the Q3 Sales sheet?",
        source="sample_sales.xlsx",
    )
    assert b["calculation"]["result"] == {"Q3 Revenue": "124000"}
    assert identities(b) == [
        ("Q3 Sales", 2, "45000"),
        ("Q3 Sales", 3, "38000"),
        ("Q3 Sales", 4, "41000"),
    ]


@pytest.mark.parametrize("word,expected", [("largest", "7000"), ("smallest", "1500")])
def test_comparison_uses_all_eligible_rows_and_exact_differences(api, word, expected):
    b = ask(
        api,
        f"Among North, South, and East, which region had the {word} revenue increase from Q3 to Q4?",
        source="sample_sales.xlsx",
    )
    assert b["calculation"]["result"] == {"1": "7000", "2": "1500", "3": "3000"}
    assert f": {expected}; matching source" in b["evidence_text"]
    assert f"{int(expected):,}" in b["answer"]
    assert {c["sheet_or_page"] for c in b["citations"]} == {"Q3 Sales"}


def test_missing_competitor_prevents_universal_highest_claim(api):
    b = ask(api, "Which region had the highest Q3 revenue?", source="sample_sales.xlsx")
    assert b["status"] == "incomplete"
    assert b["calculation"] is None
    assert b["records"][-1]["values"]["Q3 Revenue"] is None


def test_pdf_numeric_transcription_is_exact_without_model(api):
    b = ask(
        api,
        "What was the Q3 revenue for East in the report?",
        source="sample_report.pdf",
    )
    assert identities(b) == [("page_1", "table_1_row_3", "41000")]
    assert "40100" not in b["answer"]


def test_wide_row_lookup_uses_complete_original_record_not_fragment_rank(api):
    b = ask(api, "What is the renewal date for order ORD-1001 in the Wide Row sheet?")
    assert identities(b, "Renewal Date") == [("Wide Row", 2, "2027-01-15")]
    assert b["records"][0]["values"]["Order ID"] == "ORD-1001"
    assert b["records"][0]["cell_references"]["Renewal Date"] == "P2"


@pytest.mark.parametrize(
    "question",
    [
        "Q3 revenue for Atlantis?",
        "Q3 revenue for North and Atlantis?",
        "Q3 revenue for North excluding South?",
        "Q3 revenue last year?",
    ],
)
def test_unknown_or_unsupported_filter_never_silently_narrows_to_known_subset(
    api, question
):
    b = ask(api, question)
    assert b["status"] == "clarification_required"
    assert b["calculation"] is None


def test_direct_raw_workbook_oracle_confirms_blanks_and_merge_ranges(fixtures_dir):
    workbook = load_workbook(
        fixtures_dir / "sample_sales_extended.xlsx", data_only=True
    )
    assert {str(r) for r in workbook["Merged Region"].merged_cells.ranges} == {
        "A2:A3",
        "A4:A5",
    }
    sparse = workbook["Sparse Data"]
    assert not sparse.merged_cells.ranges
    assert [(sparse[f"A{r}"].value, sparse[f"B{r}"].value) for r in range(2, 7)] == [
        ("North", 45000),
        (None, 38000),
        ("East", None),
        ("West", 41000),
        (None, 22000),
    ]
    workbook.close()


def test_new_ledger_failure_rolls_back_only_new_vectors_and_records(
    api, monkeypatch, fixtures_dir, tmp_path
):
    before = api[1].service.documents()

    def fail(*args, **kwargs):
        raise RuntimeError("disk full")

    monkeypatch.setattr(api[1].service._records, "add", fail)
    data = (fixtures_dir / "sample_sales.xlsx").read_bytes()
    response = api[0].post("/documents", files={"file": ("new.xlsx", data)})
    assert response.status_code == 503
    assert not api[1].store.has_source_file("new.xlsx")
    assert api[1].service._records.get("new.xlsx") is None
    assert api[1].service.documents() == before


def test_vector_only_legacy_source_requires_reindex_without_deleting(api):
    api[1].service._records.delete("sample_sales.xlsx")
    before = api[1].store.count()
    response = api[0].post(
        "/query",
        json={"question": "North Q3 revenue", "source_file": "sample_sales.xlsx"},
    )
    assert response.status_code == 409
    assert response.json()["code"] == "reindex_required"
    assert api[1].store.count() == before


def test_narrative_is_source_excerpt_with_explicit_coverage_boundary(api):
    b = ask(
        api,
        "According to the Annual Company Overview, which areas had particularly strong momentum?",
        source="sample_report_extended.pdf",
        sheet_or_page="page_1",
    )
    assert b["status"] == "evidence_only"
    assert "customer retention" in b["answer"].lower()
    assert "new account acquisition" in " ".join(b["answer"].lower().split())
    assert b["coverage"]["complete"] is False
    assert b["calculation"] is None


def test_categorical_lookup_retains_missing_source_values_and_safe_text(api):
    b = ask(api, "Show status for order ORD-1001 in the Wide Row sheet?")
    assert b["status"] == "answered"
    assert b["records"][0]["values"]["Status"] == "Confirmed"
    assert b["records"][0]["values"]["Order ID"] == "ORD-1001"


@pytest.mark.parametrize(
    "setting,filename",
    [
        ("max_xlsx_expanded_bytes", "sample_sales.xlsx"),
        ("max_extracted_rows", "sample_sales.xlsx"),
        ("max_pdf_pages", "sample_report_extended.pdf"),
    ],
)
def test_extraction_budget_rejects_without_partial_records(
    api, monkeypatch, fixtures_dir, setting, filename
):
    from app.core.config import settings

    monkeypatch.setattr(settings, setting, 1)
    response = api[0].post(
        "/documents",
        files={"file": ("budget-" + filename, (fixtures_dir / filename).read_bytes())},
    )
    assert response.status_code == 413
    assert response.json()["code"] == "extraction_limit"
    assert not api[1].store.has_source_file("budget-" + filename)
    assert api[1].service._records.get("budget-" + filename) is None


def test_implicit_single_source_does_not_leak_stale_vector_only_documents(api):
    for source in (
        "sample_sales.xlsx",
        "sample_sales_extended.xlsx",
        "sample_report_extended.pdf",
    ):
        api[1].service._records.delete(source)
    response = api[0].post(
        "/query",
        json={
            "question": "Summarize quarterly regional sales performance",
            "top_k": 10,
        },
    )
    assert response.status_code == 200
    body = response.json()
    assert body["status"] == "evidence_only"
    assert body["citations"]
    assert {c["source_file"] for c in body["citations"]} == {"sample_report.pdf"}


def test_requested_entity_order_keeps_values_citations_and_source_rows_together(api):
    b = ask(
        api,
        "Combined Q3 revenue for East, North, and South in Q3 Sales",
        source="sample_sales.xlsx",
    )
    assert identities(b) == [
        ("Q3 Sales", 4, "41000"),
        ("Q3 Sales", 2, "45000"),
        ("Q3 Sales", 3, "38000"),
    ]
    assert b["calculation"]["result"] == {"Q3 Revenue": "124000"}
    assert [c["row_or_location"] for c in b["citations"]] == [4, 2, 3]


def test_wide_rows_with_nonunique_first_field_keep_distinct_order_identities(
    api, tmp_path
):
    from openpyxl import Workbook

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Orders"
    sheet.append(["Category", "Order ID", "Notes", "Renewal Date"])
    sheet.append(["Hardware", "INV-1001", "x" * 2000, "2027-01-15"])
    sheet.append(["Hardware", "INV-1002", "x" * 2000, "2028-02-16"])
    path = tmp_path / "orders.xlsx"
    workbook.save(path)
    workbook.close()
    assert (
        api[0]
        .post("/documents", files={"file": (path.name, path.read_bytes())})
        .status_code
        == 201
    )
    b = ask(api, "What is the renewal date for order INV-1002?", source=path.name)
    assert identities(b, "Renewal Date") == [("Orders", 3, "2028-02-16")]
    assert b["records"][0]["values"]["Order ID"] == "INV-1002"
